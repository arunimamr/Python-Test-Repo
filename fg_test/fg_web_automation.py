import os
import sys
import time
import argparse
import asyncio
import pandas as pd
from datetime import datetime
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from dotenv import load_dotenv, find_dotenv

# Load environment variables from .env
load_dotenv(find_dotenv())

try:
    from playwright.async_api import async_playwright, TimeoutError as PlaywrightTimeoutError
except ImportError:
    print("\n[ERROR] Playwright is not installed.")
    print("Please install it by running:")
    print("    pip install playwright pandas openpyxl python-dotenv")
    print("    python -m playwright install chromium\n")
    sys.exit(1)


# =====================================================================
# Authentication Helpers
# =====================================================================

async def auto_login_facilitygrid(page, username: str, password: str, timeout_ms: int = 15000) -> bool:
    """
    Attempts to automatically fill and submit login credentials on FacilityGrid login page.
    Returns True if successfully redirected away from login/signin, False otherwise.
    """
    if not username or not password:
        return False

    try:
        print(f"      -> Attempting automated login for user '{username}'...")
        await page.wait_for_timeout(2000)

        # Selectors commonly used in Angular / web login forms
        user_selector = "input[type='email'], input[name='username'], input[name='email'], input[formcontrolname='username'], input[formcontrolname='email'], input[id*='user'], input[id*='email'], input[type='text']"
        pass_selector = "input[type='password'], input[name='password'], input[formcontrolname='password'], input[id*='pass']"

        user_el = page.locator(user_selector).first
        pass_el = page.locator(pass_selector).first

        if await user_el.count() == 0 or await pass_el.count() == 0:
            try:
                await page.wait_for_selector(user_selector, timeout=5000)
                user_el = page.locator(user_selector).first
                pass_el = page.locator(pass_selector).first
            except Exception:
                print("      -> Could not locate login input fields on current page.")
                return False

        await user_el.fill(username)
        await page.wait_for_timeout(500)

        # Handle step-1 submit if password element is hidden
        if not await pass_el.is_visible():
            next_btn = page.locator("button:has-text('Next'), button:has-text('Continue'), button[type='submit'], input[type='submit']").first
            if await next_btn.count() > 0 and await next_btn.is_visible():
                await next_btn.click()
            else:
                await user_el.press("Enter")
            await page.wait_for_timeout(1500)

        try:
            await pass_el.fill(password, timeout=5000)
        except Exception:
            # Fallback force fill if still marked hidden
            await pass_el.fill(password, force=True)

        await page.wait_for_timeout(400)

        # Click submit button or press Enter
        submit_selector = "button[type='submit'], input[type='submit'], button:has-text('Sign In'), button:has-text('Log In'), button:has-text('Sign in'), button:has-text('Log in')"
        submit_btn = page.locator(submit_selector).first
        if await submit_btn.count() > 0 and await submit_btn.is_visible():
            await submit_btn.click()
        else:
            await pass_el.press("Enter")

        # Wait for potential redirect/navigation
        for _ in range(12):
            await page.wait_for_timeout(500)
            cur_url = page.url.lower()
            if "login" not in cur_url and "signin" not in cur_url:
                print("      [LOGIN SUCCESS] Successfully authenticated into FacilityGrid via credentials from .env!")
                return True

        cur_url = page.url.lower()
        if "login" not in cur_url and "signin" not in cur_url:
            print("      [LOGIN SUCCESS] Successfully authenticated into FacilityGrid via credentials from .env!")
            return True
        else:
            print("      [LOGIN NOTICE] Submitted credentials but page is still on login/auth (MFA/SSO or verification may be required).")
            return False
    except Exception as e:
        print(f"      [AUTO-LOGIN ERROR] {e}")
        return False


# =====================================================================
# DOM Extraction Helpers
# =====================================================================

def clean_text(text: str) -> str:
    if not text:
        return ""
    return " ".join(text.replace("\xa0", " ").strip().split())


async def extract_field_values_from_page(page, target_labels: list[str]) -> list[str]:
    """
    Scans the current page DOM for values corresponding to any of the given target_labels.
    Optimized to run entirely in the browser via page.evaluate to eliminate IPC overhead.
    """
    normalized_targets = [lbl.strip().lower() for lbl in target_labels]
    
    js_script = """
    (targets) => {
        const foundValues = new Set();
        
        function cleanText(text) {
            if (!text) return "";
            return text.replace(/\\xa0/g, " ").trim().replace(/\\s+/g, " ");
        }
        
        function addVal(val) {
            const cleaned = cleanText(val);
            if (cleaned && !['select', 'choose', 'none', '-- select --', '-- choose --', 'n/a', 'null', ''].includes(cleaned.toLowerCase())) {
                foundValues.add(cleaned);
            }
        }

        // 1. FacilityGrid Hidden Property Pairings: input[name='prop_nm[]'] paired with input[name='prop_val[]']
        document.querySelectorAll("input[name='prop_nm[]']").forEach(nmInp => {
            const nmVal = cleanText(nmInp.value).replace(/:/g, "").replace(/\\*/g, "").toLowerCase();
            if (targets.includes(nmVal)) {
                const row = nmInp.closest('tr');
                if (row) {
                    const valInp = row.querySelector("input[name='prop_val[]'], .saveout, label.plain-label, td:nth-child(2)");
                    if (valInp) {
                        if (valInp.tagName.toLowerCase() === 'input') {
                            addVal(valInp.value);
                        } else {
                            addVal(valInp.innerText || valInp.textContent);
                        }
                    }
                }
            }
        });

        // 2. Check all table rows across entire page
        document.querySelectorAll("tr").forEach(row => {
            let label = "";
            const th = row.querySelector("th");
            if (th) {
                label = cleanText(th.innerText || th.textContent);
            } else {
                const tdFirst = row.querySelector("td.label, td:first-child");
                if (tdFirst) {
                    label = cleanText(tdFirst.innerText || tdFirst.textContent);
                }
            }
            
            if (!label) return;
            
            const normLabel = label.replace(/:/g, "").replace(/\\*/g, "").toLowerCase();
            
            if (!targets.includes(normLabel)) {
                // Check 2-column parameter row: Col 1 = Name, Col 2 = Value
                const tds = row.querySelectorAll("td");
                if (tds.length >= 2) {
                    const col1Text = cleanText(tds[0].innerText || tds[0].textContent).replace(/:/g, "").toLowerCase();
                    if (targets.includes(col1Text)) {
                        const col2 = tds[1];
                        const inp = col2.querySelector("input:not([type='hidden']), select, textarea, label.plain-label");
                        if (inp) {
                            const tag = inp.tagName.toLowerCase();
                            if (tag === 'input' || tag === 'textarea') {
                                addVal(inp.value);
                            } else if (tag === 'select') {
                                const checkedOpt = inp.querySelector("option:checked");
                                addVal(checkedOpt ? (checkedOpt.innerText || checkedOpt.textContent) : inp.value);
                            } else {
                                addVal(inp.innerText || inp.textContent);
                            }
                        } else {
                            addVal(col2.innerText || col2.textContent);
                        }
                    }
                }
                return;
            }
            
            // Found matching label in TH / TD
            row.querySelectorAll("input:not([type='hidden'])").forEach(inp => addVal(inp.value));
            row.querySelectorAll("label, .plain-label, span.value, div.value").forEach(lbl => addVal(lbl.innerText || lbl.textContent));
            row.querySelectorAll("select").forEach(sel => {
                const checkedOpt = sel.querySelector("option:checked");
                addVal(checkedOpt ? (checkedOpt.innerText || checkedOpt.textContent) : sel.value);
            });
            row.querySelectorAll("textarea").forEach(ta => addVal(ta.value));
            
            const tdTexts = row.querySelectorAll("td:nth-child(2), td:last-child");
            tdTexts.forEach(td => addVal(td.innerText || td.textContent));
        });

        // 3. Form-group / generic labels fallback
        if (foundValues.size === 0) {
            document.querySelectorAll("label, dt, .form-label").forEach(lbl => {
                const text = cleanText(lbl.innerText || lbl.textContent).replace(/:/g, "").replace(/\\*/g, "").toLowerCase();
                if (targets.includes(text)) {
                    const parent = lbl.parentElement;
                    if (parent) {
                        parent.querySelectorAll("input:not([type='hidden']), select, textarea, span, dd, label.plain-label").forEach(el => {
                            const tag = el.tagName.toLowerCase();
                            if (tag === 'input' || tag === 'textarea') {
                                addVal(el.value);
                            } else if (tag === 'select') {
                                const checkedOpt = el.querySelector("option:checked");
                                addVal(checkedOpt ? (checkedOpt.innerText || checkedOpt.textContent) : el.value);
                            } else {
                                const val = el.innerText || el.textContent;
                                if (cleanText(val).toLowerCase() !== text) {
                                    addVal(val);
                                }
                            }
                        });
                    }
                }
            });
        }

        return Array.from(foundValues);
    }
    """
    
    try:
        return await page.evaluate(js_script, normalized_targets)
    except Exception as e:
        print(f"      [DOM EXTRACTION ERROR] {e}")
        return []


async def extract_equipment_link(page) -> tuple[str, str]:
    """
    Finds Equipment/Asset link on the Issue page or Subtask page.
    Returns (equipment_name, equipment_href).
    """
    try:
        asset_field = page.locator("#id_asset_field a, #id_asset_field")
        if await asset_field.count() > 0:
            loc_a = asset_field.locator("a")
            link = loc_a.first if await loc_a.count() > 0 else asset_field.first
            href = await link.get_attribute("href") or ""
            text = clean_text(await link.inner_text())
            if href and text and text.lower() not in ["select", "none", "-- select --"]:
                return text, href
    except Exception:
        pass

    try:
        equip_links = page.locator("a[href*='action=eequip'], a[href*='action=easset'], a[href*='action=equip'], a[href*='action=easet']")
        count = await equip_links.count()
        for i in range(count):
            link = equip_links.nth(i)
            href = await link.get_attribute("href") or ""
            text = clean_text(await link.inner_text())
            if href and text and text.lower() not in ["equipment", "asset", "select", "add"]:
                return text, href
    except Exception:
        pass

    try:
        for row_lbl in ["Equipment", "Asset", "Equipment / Asset", "Asset / Equipment", "Equipment Name", "Asset Name"]:
            row = page.locator(f"tr:has(th:text-is('{row_lbl}')), tr:has(td:text-is('{row_lbl}'))")
            if await row.count() > 0:
                link = row.first.locator("a").first
                if await link.count() > 0:
                    href = await link.get_attribute("href") or ""
                    text = clean_text(await link.inner_text())
                    if href and text and text.lower() not in ["select", "none"]:
                        return text, href
    except Exception:
        pass

    return "", ""


async def extract_subtask_link(page) -> tuple[str, str]:
    """
    Finds Subtask link on the Issue page.
    Returns (subtask_name, subtask_href).
    """
    try:
        subtask_field = page.locator("#subtask_id a, #subtask_id, [name*='subtask'] a")
        if await subtask_field.count() > 0:
            loc_a = subtask_field.locator("a")
            link = loc_a.first if await loc_a.count() > 0 else subtask_field.first
            href = await link.get_attribute("href") or ""
            text = clean_text(await link.inner_text())
            if href and text and text.lower() not in ["select", "none", "-- select --"]:
                return text, href
    except Exception:
        pass

    try:
        subtask_links = page.locator("a[href*='action=esubtask'], a[href*='action=subtask'], a[href*='action=ecst']")
        count = await subtask_links.count()
        for i in range(count):
            link = subtask_links.nth(i)
            href = await link.get_attribute("href") or ""
            text = clean_text(await link.inner_text())
            if href and text and text.lower() not in ["subtask", "select", "add"]:
                return text, href
    except Exception:
        pass

    try:
        for row_lbl in ["Subtask", "Subtask / Checklist", "Checklist Task", "Task"]:
            row = page.locator(f"tr:has(th:text-is('{row_lbl}')), tr:has(td:text-is('{row_lbl}'))")
            if await row.count() > 0:
                link = row.first.locator("a").first
                if await link.count() > 0:
                    href = await link.get_attribute("href") or ""
                    text = clean_text(await link.inner_text())
                    if href and text and text.lower() not in ["select", "none"]:
                        return text, href
    except Exception:
        pass

    return "", ""


async def extract_equipment_type_link(page) -> tuple[str, str]:
    """
    Finds Equipment Type / Asset Type link on the Equipment page.
    Returns (type_name, type_href).
    """
    type_name = ""
    try:
        lbl = page.locator("#id_eq_type_field, #id_type_field, #id_type_category_field").first
        if await lbl.count() > 0:
            type_name = clean_text(await lbl.inner_text())
    except Exception:
        pass

    try:
        edit_type = page.locator("#id_edit_type, a[href*='action=eeqt'], #id_eqt_frame a.command-sml, a[href*='action=etype']").first
        if await edit_type.count() > 0:
            href = await edit_type.get_attribute("href") or ""
            if href:
                return type_name or "Equipment Type", href
    except Exception:
        pass

    try:
        for row_lbl in ["Type", "Equipment Type", "Asset Type", "Type Name"]:
            row = page.locator(f"tr:has(th:text-is('{row_lbl}')), tr:has(td:text-is('{row_lbl}'))")
            if await row.count() > 0:
                link = row.first.locator("a").first
                if await link.count() > 0:
                    href = await link.get_attribute("href") or ""
                    text = clean_text(await link.inner_text())
                    if href and text and text.lower() not in ["select", "none"]:
                        return text or type_name, href
    except Exception:
        pass

    return type_name, ""


async def check_parameters_tab_if_present(page):
    """
    Checks if there's a 'Parameters' or 'Custom Fields' or 'Attributes' tab/link, and clicks it to load parameter values.
    """
    try:
        param_tab = page.locator("a:text-is('Parameters'), li:has-text('Parameters') a, button:has-text('Parameters'), a[href*='param']").first
        if await param_tab.count() > 0 and await param_tab.is_visible():
            await param_tab.click(timeout=3000)
            await page.wait_for_timeout(1000)
    except Exception:
        pass


def resolve_full_url(base_url: str, href: str) -> str:
    """Combines base domain with relative or inquire href."""
    if not href:
        return ""
    if href.startswith("http://") or href.startswith("https://"):
        return href
    if href.startswith("/"):
        from urllib.parse import urlparse
        parsed = urlparse(base_url)
        return f"{parsed.scheme}://{parsed.netloc}{href}"
    from urllib.parse import urljoin
    return urljoin(base_url, href)


# =====================================================================
# Multi-Level Hierarchy Inspection for a Single Issue
# =====================================================================

async def inspect_issue_and_hierarchy(page, issue_url: str, timeout_ms: int = 15000, worker_prefix: str = "") -> dict:
    """
    Executes the multi-level hierarchy flow asynchronously:
    1. Navigate to Issue URL -> check Issue for Manufacturer / Model
    2. If missing, look for Equipment/Asset -> open Equipment -> check Equipment frame, #nav_comp Parameters, & Type
    3. If missing and no Equipment, look for Subtask -> open Subtask -> check Subtask & linked Equipment
    """
    result = {
        "manufacturer": "",
        "model_number": "",
        "found_source": "",
        "equipment_name": "",
        "subtask_name": "",
        "status": "Success",
    }

    mfg_labels = ["Manufacturer", "Make", "Mfg"]
    model_labels = ["Model Number", "Model", "Model No", "Model #"]

    # --- Step 1: Open Issue URL ---
    try:
        await page.goto(issue_url, timeout=timeout_ms, wait_until="domcontentloaded")
        await page.wait_for_timeout(1000)
    except PlaywrightTimeoutError:
        result["status"] = "Timeout (Issue Page)"
        return result
    except Exception as e:
        result["status"] = f"Error (Issue Page): {str(e)[:50]}"
        return result

    # Check direct Issue page
    issue_mfgs = await extract_field_values_from_page(page, mfg_labels)
    issue_models = await extract_field_values_from_page(page, model_labels)

    if issue_mfgs:
        result["manufacturer"] = ", ".join(issue_mfgs)
    if issue_models:
        result["model_number"] = ", ".join(issue_models)

    if result["manufacturer"] and result["model_number"]:
        result["found_source"] = "Issue Page"
        return result

    # --- Step 2: Check for Equipment / Asset ---
    equip_name, equip_href = await extract_equipment_link(page)
    subtask_name, subtask_href = await extract_subtask_link(page)

    if equip_name:
        result["equipment_name"] = equip_name
    if subtask_name:
        result["subtask_name"] = subtask_name

    if equip_href:
        equip_url = resolve_full_url(issue_url, equip_href)
        if worker_prefix:
            print(f"      {worker_prefix} -> Found Equipment: '{equip_name}' | Navigating...", flush=True)
        try:
            await page.goto(equip_url, timeout=timeout_ms, wait_until="domcontentloaded")
            await page.wait_for_timeout(1500)

            # Check Equipment Level & In-page Parameters (#id_eqt_frame, #nav_comp, #eq_details)
            eq_mfgs = await extract_field_values_from_page(page, mfg_labels)
            eq_models = await extract_field_values_from_page(page, model_labels)

            if eq_mfgs and not result["manufacturer"]:
                result["manufacturer"] = ", ".join(eq_mfgs)
            if eq_models and not result["model_number"]:
                result["model_number"] = ", ".join(eq_models)

            if result["manufacturer"] or result["model_number"]:
                result["found_source"] = "Equipment Page / Parameters"

            # Check Parameters tab on Equipment if still missing
            if not (result["manufacturer"] and result["model_number"]):
                await check_parameters_tab_if_present(page)
                param_mfgs = await extract_field_values_from_page(page, mfg_labels)
                param_models = await extract_field_values_from_page(page, model_labels)
                if param_mfgs and not result["manufacturer"]:
                    result["manufacturer"] = ", ".join(param_mfgs)
                    result["found_source"] = "Equipment Parameters"
                if param_models and not result["model_number"]:
                    result["model_number"] = ", ".join(param_models)
                    result["found_source"] = "Equipment Parameters"

            # Check Equipment Type Level (a#id_edit_type or action=eeqt) if still missing
            if not (result["manufacturer"] and result["model_number"]):
                type_name, type_href = await extract_equipment_type_link(page)
                if type_href:
                    type_url = resolve_full_url(equip_url, type_href)
                    if worker_prefix:
                        print(f"      {worker_prefix} -> Found Equipment Type: '{type_name}' | Navigating...", flush=True)
                    try:
                        await page.goto(type_url, timeout=timeout_ms, wait_until="domcontentloaded")
                        await page.wait_for_timeout(1500)

                        type_mfgs = await extract_field_values_from_page(page, mfg_labels)
                        type_models = await extract_field_values_from_page(page, model_labels)

                        if type_mfgs and not result["manufacturer"]:
                            result["manufacturer"] = ", ".join(type_mfgs)
                            result["found_source"] = f"Type Page ({type_name})"
                        if type_models and not result["model_number"]:
                            result["model_number"] = ", ".join(type_models)
                            result["found_source"] = f"Type Page ({type_name})"

                        # Check Parameters tab on Type
                        if not (result["manufacturer"] and result["model_number"]):
                            await check_parameters_tab_if_present(page)
                            type_p_mfgs = await extract_field_values_from_page(page, mfg_labels)
                            type_p_models = await extract_field_values_from_page(page, model_labels)
                            if type_p_mfgs and not result["manufacturer"]:
                                result["manufacturer"] = ", ".join(type_p_mfgs)
                                result["found_source"] = f"Type Parameters ({type_name})"
                            if type_p_models and not result["model_number"]:
                                result["model_number"] = ", ".join(type_p_models)
                                result["found_source"] = f"Type Parameters ({type_name})"
                    except Exception as te:
                        if worker_prefix:
                            print(f"      {worker_prefix} [WARN] Could not inspect Type page: {str(te)[:40]}", flush=True)

        except Exception as ee:
            if worker_prefix:
                print(f"      {worker_prefix} [WARN] Could not inspect Equipment page: {str(ee)[:40]}", flush=True)

    # --- Step 3: If No Equipment or still missing, check Subtask ---
    if not (result["manufacturer"] and result["model_number"]) and subtask_href:
        subtask_url = resolve_full_url(issue_url, subtask_href)
        if worker_prefix:
            print(f"      {worker_prefix} -> Found Subtask: '{subtask_name}' | Navigating...")
        try:
            await page.goto(subtask_url, timeout=timeout_ms, wait_until="domcontentloaded")
            await page.wait_for_timeout(1500)

            # Check Subtask Level
            sub_mfgs = await extract_field_values_from_page(page, mfg_labels)
            sub_models = await extract_field_values_from_page(page, model_labels)

            if sub_mfgs and not result["manufacturer"]:
                result["manufacturer"] = ", ".join(sub_mfgs)
                result["found_source"] = "Subtask Page"
            if sub_models and not result["model_number"]:
                result["model_number"] = ", ".join(sub_models)
                result["found_source"] = "Subtask Page"

            # Check if Subtask has an Equipment link
            if not (result["manufacturer"] and result["model_number"]):
                sub_eq_name, sub_eq_href = await extract_equipment_link(page)
                if sub_eq_href:
                    sub_eq_url = resolve_full_url(subtask_url, sub_eq_href)
                    if worker_prefix:
                        print(f"      {worker_prefix} -> Subtask linked to Equipment: '{sub_eq_name}' | Navigating...")
                    try:
                        await page.goto(sub_eq_url, timeout=timeout_ms, wait_until="domcontentloaded")
                        await page.wait_for_timeout(1500)

                        s_eq_mfgs = await extract_field_values_from_page(page, mfg_labels)
                        s_eq_models = await extract_field_values_from_page(page, model_labels)
                        if s_eq_mfgs and not result["manufacturer"]:
                            result["manufacturer"] = ", ".join(s_eq_mfgs)
                            result["found_source"] = "Subtask -> Equipment"
                        if s_eq_models and not result["model_number"]:
                            result["model_number"] = ", ".join(s_eq_models)
                            result["found_source"] = "Subtask -> Equipment"
                    except Exception:
                        pass
        except Exception as se:
            if worker_prefix:
                print(f"      {worker_prefix} [WARN] Could not inspect Subtask page: {str(se)[:40]}")

    return result


# =====================================================================
# Excel Output Formatting with RED Highlighting
# =====================================================================

def save_styled_excel(records: list[dict], output_path: str):
    """
    Saves scan records to an Excel workbook with styling:
    - When BOTH Manufacturer & Model Number are found, the row & manufacturer cells are highlighted in RED!
    - Clean table formatting with column widths and borders.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Scan Results"

    headers = [
        "Excel Row",
        "Issue Code",
        "Project Name",
        "Found Source",
        "Equipment / Asset",
        "Subtask",
        "Web Manufacturer",
        "Web Model Number",
        "Both Found?",
        "Excel Manufacturer",
        "Issue URL",
        "Status",
    ]

    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")

    # High-visibility RED highlight for when BOTH Manufacturer & Model are found
    red_cell_fill = PatternFill(start_color="FF4D4D", end_color="FF4D4D", fill_type="solid")
    red_cell_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")

    # Light red row fill for both-found rows
    light_red_row_fill = PatternFill(start_color="FFE6E6", end_color="FFE6E6", fill_type="solid")
    
    # Yellow fill for partial found (only mfg or only model)
    partial_fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")

    regular_font = Font(name="Segoe UI", size=10)
    bold_font = Font(name="Segoe UI", size=10, bold=True)
    border_thin = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    # Write headers
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.row_dimensions[1].height = 28

    # Sort records by excel_row to ensure clean sequential ordering in output file
    sorted_records = sorted(records, key=lambda x: x.get("excel_row", 0))

    # Write rows
    for r_idx, r in enumerate(sorted_records, start=2):
        both_found = bool(r.get("manufacturer_found") and r.get("model_found"))
        partial_found = bool((r.get("manufacturer_found") or r.get("model_found")) and not both_found)

        row_data = [
            r.get("excel_row", ""),
            r.get("issue_code", ""),
            r.get("project_name", ""),
            r.get("found_source", ""),
            r.get("equipment_name", ""),
            r.get("subtask_name", ""),
            r.get("web_manufacturer", ""),
            r.get("web_model_number", ""),
            "YES (RED)" if both_found else ("PARTIAL" if partial_found else "NO"),
            r.get("excel_manufacturer", ""),
            r.get("issue_url", ""),
            r.get("status", ""),
        ]

        ws.row_dimensions[r_idx].height = 20

        for col_idx, val in enumerate(row_data, start=1):
            cell = ws.cell(row=r_idx, column=col_idx, value=val)
            cell.font = regular_font
            cell.border = border_thin
            cell.alignment = Alignment(vertical="center")

            if col_idx in [1, 2, 4, 9, 12]:
                cell.alignment = Alignment(horizontal="center", vertical="center")

            # Apply Row Highlighting
            if both_found:
                cell.fill = light_red_row_fill
                # Highlight Web Manufacturer, Model Number, and Both Found with RED
                if col_idx in [7, 8, 9]:
                    cell.fill = red_cell_fill
                    cell.font = red_cell_font
            elif partial_found:
                cell.fill = partial_fill
                if col_idx in [7, 8]:
                    cell.font = bold_font

    # Auto-adjust column widths
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 45)

    ws.freeze_panes = "A2"
    try:
        wb.save(output_path)
    except PermissionError:
        base, ext = os.path.splitext(output_path)
        alt_path = f"{base}_updated{ext}"
        print(f"\n[WARNING] Could not save to '{output_path}' because it is open in Excel.", flush=True)
        print(f"[INFO] Saving results to secondary output file instead: {alt_path}", flush=True)
        wb.save(alt_path)


# =====================================================================
# Parallel Worker Loop & Runner
# =====================================================================

async def worker_task(
    worker_id: int,
    context,
    queue: asyncio.Queue,
    records: list,
    records_lock: asyncio.Lock,
    total: int,
    completed_counter: list,
    batch_size: int,
    output_path: str,
    timeout_ms: int,
):
    """
    Worker task running on a dedicated Playwright page.
    Pulls rows from the queue, processes hierarchy inspection, logs live status,
    and updates batch records safely.
    """
    page = await context.new_page()
    try:
        while not queue.empty():
            try:
                item = queue.get_nowait()
            except asyncio.QueueEmpty:
                break

            idx, original_idx, row = item
            issue_code = str(row.get("issue_code", ""))
            project_name = str(row.get("project_name", ""))
            issue_url = str(row.get("issue_url", "")).strip()
            excel_mfg = str(row.get("manufacturer", "")) if pd.notna(row.get("manufacturer")) else ""

            async with records_lock:
                started_done = completed_counter[0]
                progress_pct = (started_done / total) * 100 if total > 0 else 0.0

            worker_prefix = f"[Worker-{worker_id}][{started_done}/{total} ({progress_pct:.1f}%)]"
            print(f"{worker_prefix} Checking {issue_code} ({project_name})...", flush=True)

            scan_res = await inspect_issue_and_hierarchy(page, issue_url, timeout_ms=timeout_ms, worker_prefix=worker_prefix)

            mfg_val = scan_res.get("manufacturer", "")
            model_val = scan_res.get("model_number", "")
            found_src = scan_res.get("found_source", "")
            eq_name = scan_res.get("equipment_name", "")
            sub_name = scan_res.get("subtask_name", "")
            status = scan_res.get("status", "Success")

            mfg_found = bool(mfg_val)
            model_found = bool(model_val)
            both_found = mfg_found and model_found

            rec = {
                "excel_row": original_idx + 2,
                "issue_code": issue_code,
                "project_name": project_name,
                "issue_url": issue_url,
                "excel_manufacturer": excel_mfg,
                "web_manufacturer": mfg_val,
                "web_model_number": model_val,
                "found_source": found_src,
                "equipment_name": eq_name,
                "subtask_name": sub_name,
                "manufacturer_found": mfg_found,
                "model_found": model_found,
                "status": status,
            }

            async with records_lock:
                records.append(rec)
                completed_counter[0] += 1
                curr_done = completed_counter[0]
                pct = (curr_done / total) * 100 if total > 0 else 0.0

                match_msg = " >>> [MATCH - BOTH FOUND (RED)]" if both_found else ""
                print(f"[Worker-{worker_id}][Done {curr_done}/{total} ({pct:.1f}%)] Result -> Mfg: '{mfg_val}' | Model: '{model_val}' | Source: '{found_src}'{match_msg}", flush=True)

                # Save checkpoint every batch_size completed rows or on final row
                if curr_done % batch_size == 0 or curr_done == total:
                    save_styled_excel(records, output_path)
                    print(f"      [BATCH CHECKPOINT] Saved progress ({curr_done}/{total} completed - {pct:.1f}%) to: {output_path}", flush=True)

            queue.task_done()
    finally:
        await page.close()


async def run_scanner(
    excel_path: str = None,
    output_path: str = None,
    limit: int = None,
    only_null_mfg: bool = True,
    headless: bool = False,
    timeout_ms: int = 15000,
    workers: int = 5,
    batch_size: int = 10,
):
    # Fetch from .env or fallback
    env_folder = (os.getenv("FG_FOLDER_NAME") or "").strip()
    env_file = (os.getenv("FG_WEB_TEST_FILE_NAME") or os.getenv("FG_FILE_NAME") or "").strip()
    fg_username = (os.getenv("FG_USERNAME") or os.getenv("FG_EMAIL") or "").strip()
    fg_password = (os.getenv("FG_PASSWORD") or "").strip()

    # Determine worker count from .env if not explicitly set via CLI
    if workers is None or workers <= 0:
        env_workers = os.getenv("FG_WEB_WORKERS") or os.getenv("WEB_WORKERS") or os.getenv("CONCURRENCY")
        try:
            workers = int(env_workers) if env_workers else 5
        except ValueError:
            workers = 5

    # Resolve input Excel file path
    if not excel_path:
        if env_folder and env_file:
            excel_path = os.path.join(env_folder, env_file)
        elif env_file:
            excel_path = env_file
        else:
            print("[ERROR] Please provide --file or configure FG_FOLDER_NAME & FG_WEB_TEST_FILE_NAME in .env")
            return
    elif env_folder and not os.path.isabs(excel_path) and not os.path.exists(excel_path):
        candidate = os.path.join(env_folder, excel_path)
        if os.path.exists(candidate):
            excel_path = candidate

    if not os.path.exists(excel_path):
        print(f"\n[ERROR] Excel input file not found: {excel_path}")
        if env_folder:
            print(f"      (Folder checked: {env_folder})")
        print("Please check FG_FOLDER_NAME and FG_FILE_NAME in your .env file, or pass --file explicitly.\n")
        return

    # Determine output folder: prefer env_folder, or input file folder
    target_folder = env_folder if env_folder and os.path.isdir(env_folder) else os.path.dirname(os.path.abspath(excel_path))
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")

    # Resolve output Excel file path
    if not output_path:
        default_name = f"scan_results_{timestamp}.xlsx"
        output_path = os.path.join(target_folder, default_name) if target_folder else default_name
    else:
        if not os.path.isabs(output_path) and not os.path.dirname(output_path) and target_folder:
            output_path = os.path.join(target_folder, output_path)

    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)

    print(f"\n=======================================================")
    print(f"Loading Excel file  : {os.path.abspath(excel_path)}")
    print(f"Output destination  : {os.path.abspath(output_path)}")
    print(f"Parallel Workers    : {workers}")
    print(f"Batch Checkpoint Size: {batch_size}")
    df = pd.read_excel(excel_path)
    print(f"Total rows in Excel : {len(df)}")

    valid_df = df[df["issue_url"].notna() & (df["issue_url"].astype(str).str.strip() != "")].copy()
    print(f"Rows with issue_url : {len(valid_df)}")

    if only_null_mfg and "manufacturer" in valid_df.columns:
        valid_df = valid_df[valid_df["manufacturer"].isna() | (valid_df["manufacturer"].astype(str).str.strip() == "")]
        print(f"Rows with NULL/empty manufacturer: {len(valid_df)}")

    if limit and limit > 0:
        valid_df = valid_df.head(limit)
        print(f"Applying limit     : Processing top {limit} rows")

    total_rows = len(valid_df)
    if total_rows == 0:
        print("[INFO] No rows to process matching the criteria.")
        return

    profile_dir = os.path.join(os.path.dirname(os.path.abspath(__file__)), ".fg_playwright_profile")
    os.makedirs(profile_dir, exist_ok=True)
    print(f"Using persistent browser profile: {profile_dir}")

    async with async_playwright() as p:
        context = await p.chromium.launch_persistent_context(
            user_data_dir=profile_dir,
            headless=headless,
            args=["--start-maximized", "--disable-blink-features=AutomationControlled"],
            no_viewport=True,
            user_agent="Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36",
        )

        init_page = context.pages[0] if context.pages else await context.new_page()

        # One-time login check on the first issue URL
        first_row = valid_df.iloc[0]
        initial_url = str(first_row.get("issue_url", "")).strip()
        if initial_url:
            print(f"\n[INIT] Checking authentication status at: {initial_url} ...")
            try:
                await init_page.goto(initial_url, timeout=timeout_ms, wait_until="domcontentloaded")
                await init_page.wait_for_timeout(2000)
                if "login" in init_page.url.lower() or "signin" in init_page.url.lower():
                    logged_in = False
                    if fg_username and fg_password:
                        print(f"\n[INIT] Login page detected. Attempting automated login using credentials from .env...")
                        logged_in = await auto_login_facilitygrid(init_page, fg_username, fg_password, timeout_ms=timeout_ms)

                    if not logged_in:
                        if headless:
                            print("\n[WARNING] Login required, but running in headless mode! Please run without --headless first to log in.")
                        print("\n" + "=" * 60)
                        print(">>> [ACTION REQUIRED] Please complete login on the opened Chromium browser.")
                        print(">>> Once you are logged in and can see FacilityGrid, press ENTER here.")
                        print("=" * 60 + "\n")
                        await asyncio.to_thread(input, ">>> Press ENTER here after you have logged in to continue automation... ")
                else:
                    print("[INIT] Active login session detected! Proceeding directly with parallel batch scanning...\n")
            except Exception as e:
                print(f"[INIT] Note: {e}")

        # Populate Queue
        queue = asyncio.Queue()
        for idx, (original_idx, row) in enumerate(valid_df.iterrows(), start=1):
            queue.put_nowait((idx, original_idx, row))

        records = []
        records_lock = asyncio.Lock()
        completed_counter = [0]

        # Launch parallel worker tasks
        num_workers = min(workers, total_rows)
        print(f"Starting {num_workers} parallel browser worker pages...\n")
        worker_tasks = [
            asyncio.create_task(
                worker_task(
                    worker_id=w_idx + 1,
                    context=context,
                    queue=queue,
                    records=records,
                    records_lock=records_lock,
                    total=total_rows,
                    completed_counter=completed_counter,
                    batch_size=batch_size,
                    output_path=output_path,
                    timeout_ms=timeout_ms,
                )
            )
            for w_idx in range(num_workers)
        ]

        await asyncio.gather(*worker_tasks)
        await context.close()

    # Final save & statistics
    save_styled_excel(records, output_path)

    out_df = pd.DataFrame(records)
    mfg_count = out_df["manufacturer_found"].sum() if "manufacturer_found" in out_df else 0
    model_count = out_df["model_found"].sum() if "model_found" in out_df else 0
    both_count = (out_df["manufacturer_found"] & out_df["model_found"]).sum() if "manufacturer_found" in out_df and "model_found" in out_df else 0

    print(f"\n==================== SCAN COMPLETE ====================")
    print(f"Total Rows Checked       : {len(out_df)}")
    print(f"Rows with Manufacturer   : {mfg_count}")
    print(f"Rows with Model Number   : {model_count}")
    print(f"Rows with BOTH Found     : {both_count} (Highlighted in RED)")
    print(f"Results saved to         : {os.path.abspath(output_path)}")
    print(f"=======================================================\n")


def main():
    parser = argparse.ArgumentParser(description="Scan Issue URLs & Hierarchy for Manufacturer & Model Number in Parallel Batches")
    parser.add_argument(
        "--file",
        "-f",
        default=None,
        help="Path to the Excel file (optional, defaults to FG_FOLDER_NAME and FG_WEB_TEST_FILE_NAME from .env)",
    )
    parser.add_argument("--output", "-o", default=None, help="Output Excel file path (defaults to same folder as input)")
    parser.add_argument("--limit", "-l", type=int, default=10, help="Number of rows to test (default: 10, set 0 for all)")
    parser.add_argument("--all-rows", action="store_true", help="Test all rows including those with existing manufacturer in Excel")
    parser.add_argument("--headless", action="store_true", help="Run browser in headless mode (default: visible UI)")
    parser.add_argument("--timeout", type=int, default=15000, help="Timeout in milliseconds per page (default: 15000)")
    parser.add_argument("--workers", "-w", type=int, default=5, help="Number of parallel worker browser pages (default: 5)")
    parser.add_argument("--batch-size", "-b", type=int, default=10, help="Excel checkpoint save frequency (default: 10 rows)")

    args = parser.parse_args()

    asyncio.run(
        run_scanner(
            excel_path=args.file,
            output_path=args.output,
            limit=args.limit if args.limit > 0 else None,
            only_null_mfg=not args.all_rows,
            headless=args.headless,
            timeout_ms=args.timeout,
            workers=args.workers,
            batch_size=args.batch_size,
        )
    )


if __name__ == "__main__":
    main()
